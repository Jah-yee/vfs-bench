import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parents[3] / "server" / "clients"))

import gws  # noqa: E402

FOLDER_NAME = "Music-Analysis"
SOURCE_SHEET = "Billboard Pop Chart by Year"
# What both arms are told the folder is; see the note in
# nhl-b2b-analysis/shared/seed_mock.py. This task's evaluator does not read
# the id back -- it finds the spreadsheet by name -- but `DRIVE_FOLDER_ID`
# still scopes the MCP server and the `gws` CLI, so a drift would put the
# agent's work somewhere neither arm is looking.
FOLDER_ID = "f0001"
REPO_ROOT = HERE.parents[3]
WORKBOOK = (REPO_ROOT / "submodule" / "Toolathlon" / "tasks" / "finalpool" /
            "music-analysis" / "groundtruth_workspace" /
            "Billboard Pop Chart by Year.xlsx")


def tabs_from_workbook() -> list[tuple[str, list[list[str]]]]:
    """Read every worksheet as rows of strings.

    Returns:
        list[tuple[str, list[list[str]]]]: (tab title, rows) in file order.
    """
    import openpyxl
    book = openpyxl.load_workbook(WORKBOOK, read_only=True, data_only=True)
    out = []
    for title in book.sheetnames:
        rows = [["" if cell is None else str(cell) for cell in row]
                for row in book[title].iter_rows(values_only=True)]
        while rows and not any(cell for cell in rows[-1]):
            rows.pop()
        out.append((title, rows))
    book.close()
    return out


def main() -> None:
    """Reset the fake server and lay down the folder, sheet and tabs."""
    if not WORKBOOK.is_file():
        sys.exit(f"missing {WORKBOOK}; init the Toolathlon submodule")

    tabs = tabs_from_workbook()
    gws.reset()
    folder_id = gws.create_folder(FOLDER_NAME)
    if folder_id != FOLDER_ID:
        raise SystemExit(f"the mock gave the folder id {folder_id!r}, "
                         f"but both arms are configured for {FOLDER_ID!r}")
    sheet_id = gws.create_sheet(SOURCE_SHEET, folder_id)

    # The first tab already exists; the rest are added. Each is declared at
    # its own dimensions so the grid matches the source workbook exactly.
    # `values.update` would grow it anyway -- the mock auto-expands past the
    # 1000x26 default the same way Sheets does -- but then every tab would
    # report whatever the last write happened to reach.
    first_title, _ = tabs[0]
    widths = {
        title: max((len(r) for r in rows), default=1)
        for title, rows in tabs
    }
    heights = {title: max(len(rows), 1) for title, rows in tabs}
    requests: list[dict] = [{
        "updateSheetProperties": {
            "properties": {
                "sheetId": 0,
                "title": first_title,
                "gridProperties": {
                    "rowCount": heights[first_title],
                    "columnCount": widths[first_title],
                },
            },
            "fields": "title,gridProperties",
        }
    }]
    for title, _ in tabs[1:]:
        requests.append({
            "addSheet": {
                "properties": {
                    "title": title,
                    "gridProperties": {
                        "rowCount": heights[title],
                        "columnCount": widths[title],
                    },
                }
            }
        })
    gws.batch_update(sheet_id, requests)

    total = 0
    for title, rows in tabs:
        if not rows:
            continue
        # Tab titles carry spaces and commas ("Top 1,000"); the client
        # escapes the A1 range before it goes into the path.
        gws.write_values(sheet_id, f"{title}!A1", rows)
        total += len(rows)

    print(f"GWS_URL={gws.base()}")
    print(f"folder_id={folder_id}")
    print(f"source_spreadsheet_id={sheet_id}")
    print(f"tabs={len(tabs)} rows={total}")


if __name__ == "__main__":
    sys.exit(main())

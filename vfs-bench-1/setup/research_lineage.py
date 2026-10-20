"""Check research conversions against their original sources."""
import hashlib
import html
import json
import re
from pathlib import Path

import openpyxl
import yaml

from server.fakes.sheet_formats import formatted
from server.seed.browser import TASKS as BROWSER_TASKS
from server.seed.business_sheets import cell_format, source_folder
from server.seed.kubernetes import original
from server.seed.papers import TASKS as PAPER_TASKS
from server.seed.research import TASKS
from server.seed.research_sheets import TASKS as SHEET_TASKS

REPO = Path(__file__).resolve().parents[1]


def main():
    configs = {
        'arxiv_local': 'arxiv_local',
        'scholarly': 'scholarly_search',
        'fetch': 'npx-fetch'
    }
    for task in TASKS:
        prefix = f'tasks/finalpool/{task}/'
        folder = REPO / 'task/toolathlon' / task
        assert (
            folder /
            'shared/task.md').read_bytes() == original(prefix +
                                                       'docs/task.md'), task
        config = json.loads(original(prefix + 'task_config.json'))
        servers = config['needed_mcp_servers']
        workspace = yaml.safe_load((folder / 'vfs/workspace.yaml').read_text())
        clis = workspace['clis']
        assert ('web-search' in clis) == ('web_search'
                                          in config['needed_local_tools'])
        assert ('browser' in clis) == ('playwright_with_chunk' in servers)
        assert ('fetch' in clis) == ('fetch' in servers)
        assert ('arxiv-local' in clis) == ('arxiv_local' in servers)
        assert ('scholarly' in clis) == ('scholarly' in servers)
        assert 'gws' not in clis
        assert ('google-sheet' in clis) == ('google_sheet' in servers)
        if 'google-sheet' in clis:
            assert task in SHEET_TASKS
            assert clis['google-sheet']['cli'] == (
                '../../../../cli/toolathlon/google.mjs:SHEETS')
            assert clis['google-sheet']['config']['apiBase'] == (
                '${SHEETS_CLI_URL}')
            assert workspace['mounts']['/gd']['config'][
                'folderId'] == '${DRIVE_FOLDER_ID}'
            folder = source_folder(task)
            capture = json.loads((folder / 'source.json').read_text())
            assert hashlib.sha256(original(
                capture['source'])).hexdigest() == capture['source_sha256']
            assert hashlib.sha256(
                (folder /
                 'source.xlsx').read_bytes()).hexdigest() == capture['sha256']
            if 'title_source' in capture:
                page = (folder / 'source.html').read_bytes()
                assert hashlib.sha256(
                    page).hexdigest() == capture['title_source']['sha256']
                title = re.search(r'<title>(.*?)</title>', page.decode())[1]
                assert html.unescape(
                    title.removesuffix(' - Google Sheets')) == capture['title']
            with (folder / 'source.xlsx').open('rb') as stream:
                workbook = openpyxl.load_workbook(stream)
                assert workbook.sheetnames == capture['tabs']
                for tab in workbook:
                    for row in tab:
                        for cell in row:
                            cell_format(cell)
                rendering = folder / 'rendering.json'
                if rendering.exists():
                    for tab, record in json.loads(
                            rendering.read_text()).items():
                        body = record['body']
                        assert hashlib.sha256(
                            body.encode()).hexdigest() == record['sha256']
                        table = json.loads(
                            body[body.index('{'):body.rindex('}') +
                                 1])['table']
                        for y, row in enumerate(table['rows']):
                            for x, cell in enumerate(row['c']):
                                if not cell or not isinstance(
                                        cell.get('v'), (int, float)):
                                    continue
                                if table['cols'][x].get(
                                        'pattern') == 'General':
                                    continue
                                source = workbook[tab].cell(y + 1, x + 1)
                                assert source.value == cell['v']
                                assert formatted(
                                    source.value,
                                    cell_format(source)) == cell['f']
                workbook.close()
        if 'browser' in clis:
            assert task in BROWSER_TASKS
        assert PAPER_TASKS.get(task, []) == [
            name for name in servers if name in ('arxiv_local', 'scholarly')
        ]
        for cli, source in configs.items():
            if cli in clis:
                upstream = yaml.safe_load(
                    original('configs/mcp_servers/' + source + '.yaml'))
                assert clis[cli]['config']['timeoutMs'] == (
                    1000 * upstream['client_session_timeout_seconds'])
        # Both arms seed and grade from these original bytes at run time.
        source = REPO / 'submodule/Toolathlon' / prefix
        for directory in ('initial_workspace', 'evaluation', 'preprocess',
                          'groundtruth_workspace'):
            for path in (source / directory).rglob('*'):
                if path.is_file() and '__pycache__' not in path.parts:
                    assert path.read_bytes() == original(
                        prefix + str(path.relative_to(source))), str(path)
    print(
        f'{len(TASKS)} research prompts, files, graders and CLI mappings match'
    )


if __name__ == '__main__':
    main()

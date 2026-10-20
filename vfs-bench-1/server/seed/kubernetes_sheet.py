"""Seed the captured original audit workbook through the Sheets API."""
import json
import os
from contextlib import contextmanager
from datetime import datetime, timedelta, timezone
from pathlib import Path

import openpyxl

from server.clients import gws
from server.fakes.sheet_formats import serve as formats

REPO = Path(__file__).resolve().parents[2]


@contextmanager
def serve():
    previous = os.environ.get('GWS_URL')
    with formats(gws.base()) as endpoint:
        os.environ['GWS_URL'] = endpoint
        try:
            with seed() as record:
                yield record
        finally:
            if previous is None:
                os.environ.pop('GWS_URL', None)
            else:
                os.environ['GWS_URL'] = previous


@contextmanager
def seed():
    source = (REPO / 'task/toolathlon/k8s-safety-audit' /
              'shared/fixtures/audit.xlsx')
    workbook = openpyxl.load_workbook(source, data_only=False)
    gws.reset()
    folder = gws.create_folder('K8sSafetyAudit')
    spreadsheet = gws.create_sheet('Kubernetes Security Audit', folder)
    tabs = gws.sheets(spreadsheet)
    requests = [{
        'updateSheetProperties': {
            'properties': {
                'sheetId': tabs[0]['properties']['sheetId'],
                'title': workbook.sheetnames[0]
            },
            'fields': 'title'
        }
    }]
    requests += [{
        'addSheet': {
            'properties': {
                'title': name
            }
        }
    } for name in workbook.sheetnames[1:]]
    gws.batch_update(spreadsheet, requests)
    properties = gws.sheets(spreadsheet)
    ids = {
        tab['properties']['title']: tab['properties']['sheetId']
        for tab in properties
    }
    for tab in workbook:
        rows = [[value if value is not None else '' for value in row]
                for row in tab.values]
        gws.write_values(spreadsheet, f"'{tab.title}'!A1", rows)
        formats = []
        for row in tab:
            cells = []
            for cell in row:
                fmt = {
                    'textFormat': {
                        'fontFamily':
                        cell.font.name,
                        **({
                            'fontSize': int(cell.font.sz)
                        } if cell.font.sz else {}), 'bold':
                        bool(cell.font.b),
                        'italic':
                        bool(cell.font.i)
                    }
                }
                borders = {}
                for edge in ('top', 'bottom', 'left', 'right'):
                    border = getattr(cell.border, edge)
                    if border is None or border.style is None:
                        continue
                    # The captured workbook uses thin black header borders.
                    # Fail if a later capture needs a new format mapping.
                    if (border.style != 'thin' or border.color.type != 'rgb'
                            or border.color.rgb != 'FF000000'):
                        raise ValueError('Unsupported captured Sheet border')
                    borders[edge] = {
                        'style': 'SOLID',
                        'color': {
                            'red': 0,
                            'green': 0,
                            'blue': 0
                        }
                    }
                if borders:
                    fmt['borders'] = borders
                if cell.alignment.horizontal:
                    alignment = cell.alignment.horizontal.upper()
                    fmt['horizontalAlignment'] = alignment
                if cell.alignment.vertical:
                    fmt['verticalAlignment'] = cell.alignment.vertical.upper()
                cells.append({'userEnteredFormat': fmt})
            formats.append({'values': cells})
        gws.batch_update(spreadsheet, [{
            'updateCells': {
                'start': {
                    'sheetId': ids[tab.title],
                    'rowIndex': 0,
                    'columnIndex': 0
                },
                'rows': formats,
                'fields': 'userEnteredFormat'
            }
        }])
    workbook.close()
    target = (REPO / 'submodule/Toolathlon/tasks/finalpool' /
              'k8s-safety-audit/files/folder_id.txt')
    previous = target.read_bytes() if target.exists() else None
    credentials = REPO / 'submodule/Toolathlon/configs/google_credentials.json'
    old_credentials = credentials.read_bytes() if credentials.exists(
    ) else None
    try:
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(folder)
        credentials.write_text(
            json.dumps({
                'client_id':
                'fixture',
                'client_secret':
                'fixture',
                'refresh_token':
                gws.TOKEN,
                'token':
                gws.TOKEN,
                'expiry':
                (datetime.now(timezone.utc) +
                 timedelta(days=1)).isoformat().replace('+00:00', 'Z')
            }))
        yield {'folder_id': folder, 'sheet_id': spreadsheet}
    finally:
        if old_credentials is None:
            credentials.unlink(missing_ok=True)
        else:
            credentials.write_bytes(old_credentials)
        if previous is None:
            target.unlink(missing_ok=True)
        else:
            target.write_bytes(previous)
